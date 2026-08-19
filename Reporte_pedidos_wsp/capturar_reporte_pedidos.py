"""
capturar_reporte_pedidos.py

Captura cada tabla de supervisores del Excel de pedidos como PNG y las envía
a WhatsApp. Usa mutex de archivo compartido en C:\\proyectos\\locks\\ para
no conflictuar con Reportes_ssff_wsp/ u otros proyectos que usen Excel/COM.

Uso:
    python capturar_reporte_pedidos.py --hora 9 --destino canal
    python capturar_reporte_pedidos.py --hora 9 --destino test
    python capturar_reporte_pedidos.py --hora 9 --solo-imagenes
"""

import argparse
import ctypes
import json
import logging
import os
import random
import sys
import time
from datetime import datetime, date
from pathlib import Path

import openpyxl
import xlwings as xw
from openpyxl.utils import get_column_letter

# Importar WhatsAppClient del proyecto hermano
REPORTES_DIR = Path(__file__).parent.parent / 'Reportes_ssff_wsp'
sys.path.insert(0, str(REPORTES_DIR))
from wa_client import WhatsAppClient

# ══════════════════════════════════════════════════════════════════════════════
# CONFIGURACIÓN
# ══════════════════════════════════════════════════════════════════════════════

SCRIPT_DIR   = Path(__file__).parent.absolute()
CONFIG_PATH  = REPORTES_DIR / 'config.json'
IMAGENES_DIR = SCRIPT_DIR / 'imagenes'
IMAGENES_DIR.mkdir(exist_ok=True)

LOCKS_DIR         = Path('C:/proyectos/locks')
LOCKS_DIR.mkdir(parents=True, exist_ok=True)
GLOBAL_LOCK_NAME  = 'EXCEL_COM'
STALE_SECONDS     = 300

HORA_LBL = {
    8: '8AM', 9: '9AM', 10: '10AM', 11: '11AM', 12: '12PM',
    13: '1PM', 14: '2PM', 15: '3PM', 16: '4PM', 17: '5PM', 18: '6PM'
}

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
_logger = logging.getLogger('capturar_pedidos')


# ══════════════════════════════════════════════════════════════════════════════
# MUTEX DE ARCHIVO (misma lógica que screenshot_safe.py)
# ══════════════════════════════════════════════════════════════════════════════

def _adquirir_lock(lock_file: Path, identificador: str, timeout: int = 120) -> bool:
    inicio = time.time()
    while time.time() - inicio < timeout:
        try:
            fd = os.open(str(lock_file), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            with os.fdopen(fd, 'w') as f:
                f.write(f'PID={os.getpid()}\nSTART={datetime.now().isoformat()}\nID={identificador}\n')
            return True
        except FileExistsError:
            try:
                edad = int(time.time() - lock_file.stat().st_mtime)
            except Exception:
                edad = 0
            if edad > STALE_SECONDS:
                _logger.warning(f'[LOCK] Lock stale ({edad}s): {lock_file.name} — limpiando')
                try:
                    lock_file.unlink()
                    continue
                except Exception:
                    pass
            time.sleep(0.5)
        except Exception as e:
            _logger.error(f'[LOCK] Error: {e}')
            return False
    _logger.error(f'[LOCK] Timeout ({timeout}s) esperando {lock_file.name}')
    return False


def _liberar_lock(lock_file: Path):
    try:
        lock_file.unlink(missing_ok=True)
    except Exception as e:
        _logger.error(f'[LOCK] Error liberando {lock_file.name}: {e}')


# ══════════════════════════════════════════════════════════════════════════════
# CAPTURA VIA EXCEL COM
# ══════════════════════════════════════════════════════════════════════════════

def _capturar_rango(worksheet, rango: str, ruta_png: str, escala: float = 2.5) -> bool:
    """Captura un rango de Excel como imagen PNG via COM + clipboard."""
    try:
        from PIL import ImageGrab, Image
        import win32gui
    except ImportError:
        _logger.error('[ERROR] Requeridas: PIL, pywin32 (pip install pillow pywin32)')
        return False

    try:
        xl_range = worksheet.range(rango)

        # Traer Excel al frente
        try:
            app = worksheet.book.app
            app.api.Visible = True
            hwnd = [0]

            def _cb(h, _):
                if win32gui.IsWindowVisible(h) and 'Microsoft Excel' in win32gui.GetWindowText(h):
                    hwnd[0] = h
                    return False
                return True

            win32gui.EnumWindows(_cb, None)
            if hwnd[0]:
                ctypes.windll.user32.AllowSetForegroundWindow(
                    ctypes.windll.kernel32.GetCurrentProcessId()
                )
                win32gui.ShowWindow(hwnd[0], 9)
                win32gui.SetForegroundWindow(hwnd[0])
            time.sleep(1.0)
        except Exception as e:
            _logger.warning(f'[WARN] No se pudo traer Excel al frente: {e}')

        for intento in range(3):
            try:
                xl_range.api.CopyPicture(Appearance=1, Format=2)
                time.sleep(1.0)
                break
            except Exception as e:
                if intento < 2:
                    _logger.warning(f'[WARN] CopyPicture intento {intento+1}: {e}')
                    time.sleep(1 + intento * 0.5)
                else:
                    raise

        img = ImageGrab.grabclipboard()
        if not img:
            _logger.error('[ERROR] Clipboard vacío')
            return False

        ancho = int(img.width * escala)
        alto  = int(img.height * escala)
        img   = img.resize((ancho, alto), resample=Image.LANCZOS)
        Path(ruta_png).parent.mkdir(parents=True, exist_ok=True)
        img.save(ruta_png, 'PNG')
        _logger.info(f'[OK] {ruta_png} ({ancho}x{alto})')
        return True

    except Exception as e:
        _logger.error(f'[ERROR] capturar_rango: {e}')
        return False


# ══════════════════════════════════════════════════════════════════════════════
# DETECCIÓN DE TABLAS
# ══════════════════════════════════════════════════════════════════════════════

def detectar_tablas(archivo_excel: str) -> list:
    """
    Detecta rangos de tablas de supervisor en la hoja PEDIDOS.
    Encabezado supervisor: celda A con fondo FF255663.
    Fin de tabla: fila con valor 'TOTAL' en columna A.

    Retorna [(nombre_supervisor, rango_excel), ...]
    """
    try:
        wb = openpyxl.load_workbook(archivo_excel, data_only=True)
        if 'PEDIDOS' not in wb.sheetnames:
            _logger.error('[ERROR] Hoja PEDIDOS no encontrada')
            return []
        ws = wb['PEDIDOS']
        max_col = ws.max_column
        max_row = ws.max_row
        col_fin = get_column_letter(max_col)

        bloques = []
        i = 1
        while i <= max_row:
            c = ws.cell(i, 1)
            fg = c.fill.fgColor
            color = fg.rgb if fg.type == 'rgb' else ''
            if color.upper() == 'FF255663' and c.value and str(c.value).strip() != 'TOTAL':
                nombre = str(c.value).strip()
                fila_ini = i
                fila_fin = fila_ini
                for j in range(i + 1, min(i + 50, max_row + 1)):
                    v = ws.cell(j, 1).value
                    if v and str(v).strip() == 'TOTAL':
                        fila_fin = j
                        break
                if fila_fin > fila_ini:
                    rango = f'A{fila_ini}:{col_fin}{fila_fin}'
                    bloques.append((nombre, rango))
                    i = fila_fin + 1
                else:
                    i += 1
            else:
                i += 1

        _logger.info(f'[OK] {len(bloques)} tablas detectadas')
        return bloques
    except Exception as e:
        _logger.error(f'[ERROR] detectar_tablas: {e}')
        return []


# ══════════════════════════════════════════════════════════════════════════════
# CAPTURA PRINCIPAL
# ══════════════════════════════════════════════════════════════════════════════

def capturar_tablas(archivo_excel: str, hora: int) -> list:
    """
    Captura cada tabla de supervisor como PNG usando el mutex global de Excel.
    Retorna [(nombre_supervisor, ruta_png), ...]
    """
    etiqueta  = HORA_LBL.get(hora, f'{hora}H')
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    tablas    = detectar_tablas(archivo_excel)

    if not tablas:
        return []

    lock_global    = LOCKS_DIR / f'{GLOBAL_LOCK_NAME}.lock'
    lock_identidad = LOCKS_DIR / f'SSFF_Pedidos_{etiqueta}.lock'

    _logger.info(f'[LOCK] Solicitando EXCEL_COM...')
    if not _adquirir_lock(lock_global, f'SSFF_Pedidos_{etiqueta}', timeout=60):
        return []
    _adquirir_lock(lock_identidad, f'SSFF_Pedidos_{etiqueta}', timeout=5)
    _logger.info(f'[LOCK] EXCEL_COM adquirido')

    resultado = []
    try:
        app   = xw.App(visible=True)
        libro = app.books.open(archivo_excel)
        hoja  = libro.sheets['PEDIDOS']

        for nombre, rango in tablas:
            slug    = nombre.replace(' ', '_')[:25]
            ruta_png = str(IMAGENES_DIR / f'PEDIDOS_{slug}_{etiqueta}_{timestamp}.png')
            _logger.info(f'[CAPTURE] {nombre} ({rango})...')
            if _capturar_rango(hoja, rango, ruta_png, escala=2.5):
                resultado.append((nombre, ruta_png))
            else:
                _logger.error(f'[ERROR] Fallo captura {nombre}')
            time.sleep(0.5)

        libro.close()
        app.quit()
    except Exception as e:
        _logger.error(f'[ERROR] Error durante captura: {e}')
    finally:
        _liberar_lock(lock_identidad)
        _liberar_lock(lock_global)
        _logger.info('[LOCK] EXCEL_COM liberado')

    return resultado


# ══════════════════════════════════════════════════════════════════════════════
# ENVÍO WHATSAPP
# ══════════════════════════════════════════════════════════════════════════════

def cargar_config() -> dict:
    with open(CONFIG_PATH, 'r', encoding='utf-8') as f:
        return json.load(f)


def enviar_a_whatsapp(capturas: list, destino: str, hora: int, config: dict) -> bool:
    if not capturas:
        _logger.error('[ERROR] No hay imágenes para enviar')
        return False

    etiqueta = HORA_LBL.get(hora, f'{hora}H')
    destinos = config.get('cortes_horarios', {}).get('destinos', {})
    numero   = destinos.get(destino)
    if not numero:
        _logger.error(f'[ERROR] Destino "{destino}" no configurado en config.json')
        return False

    # Grupos individuales por supervisor (mismo dict que usa capturar_cortes.py)
    vendedor_grupos      = config.get('vendedor_grupos', {})
    vendedor_grupos_norm = {k.strip().upper(): v for k, v in vendedor_grupos.items()}

    wa     = WhatsAppClient(config_path=str(CONFIG_PATH))
    exitos = 0

    DIAS_ES = {0: 'Lunes', 1: 'Martes', 2: 'Miércoles', 3: 'Jueves',
               4: 'Viernes', 5: 'Sábado', 6: 'Domingo'}
    hoy = datetime.now()
    nombre_dia = DIAS_ES[hoy.weekday()]
    fecha_str  = hoy.strftime('%d/%m')

    for nombre_sup, ruta_png in capturas:
        if not Path(ruta_png).exists():
            _logger.warning(f'[WARN] Imagen no encontrada: {ruta_png}')
            continue
        caption = f'Corte Pedidos {nombre_dia} {fecha_str} - {etiqueta} | {nombre_sup}'

        # Enviar al canal general
        _logger.info(f'[SEND] {nombre_sup} → canal general...')
        r = wa.send_image(numero, ruta_png, caption=caption)
        if r.get('success'):
            _logger.info(f'[OK] {nombre_sup} enviado al canal')
            exitos += 1
        else:
            _logger.error(f'[ERROR] {nombre_sup}: {r.get("error")}')
        time.sleep(random.uniform(3, 6))

        # Enviar al grupo individual del supervisor si está configurado
        grupo_sup = vendedor_grupos_norm.get(nombre_sup.strip().upper())
        if grupo_sup:
            _logger.info(f'[SEND] {nombre_sup} → grupo individual...')
            r2 = wa.send_image(grupo_sup, ruta_png, caption=caption)
            if r2.get('success'):
                _logger.info(f'[OK] {nombre_sup} enviado a grupo individual')
            else:
                _logger.error(f'[ERROR] Grupo individual {nombre_sup}: {r2.get("error")}')
            time.sleep(random.uniform(3, 6))

    _logger.info(f'[OK] {exitos}/{len(capturas)} imágenes enviadas a {destino}')
    return exitos > 0


# ══════════════════════════════════════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(description='Captura y envía reporte de pedidos a WhatsApp')
    parser.add_argument('--hora',          type=int, required=True)
    parser.add_argument('--destino',       default='test')
    parser.add_argument('--archivo',       default=None)
    parser.add_argument('--solo-imagenes', action='store_true')
    args = parser.parse_args()

    etiqueta = HORA_LBL.get(args.hora, f'{args.hora}H')

    if not args.archivo:
        hoy = date.today().strftime('%Y%m%d')
        args.archivo = str(SCRIPT_DIR / f'CORTE_PEDIDOS_{hoy}.xlsx')
        _logger.info(f'[FILE] Usando: {args.archivo}')

    if not Path(args.archivo).exists():
        _logger.error(f'[ERROR] Archivo no encontrado: {args.archivo}')
        sys.exit(1)

    config   = cargar_config()
    capturas = capturar_tablas(args.archivo, args.hora)

    if not capturas:
        _logger.error('[ERROR] No se capturó ninguna imagen')
        sys.exit(1)

    if args.solo_imagenes:
        _logger.info(f'[OK] {len(capturas)} imágenes capturadas (--solo-imagenes)')
        return

    if not enviar_a_whatsapp(capturas, args.destino, args.hora, config):
        sys.exit(1)

    _logger.info(f'[OK] REPORTE PEDIDOS {etiqueta} COMPLETADO')


if __name__ == '__main__':
    main()
