"""
generar_estacional.py
---------------------
Análisis de estacionalidad por categoría para apoyo en fijación de cuotas.

Lógica:
  - Índice estacional = ventas_mes_i / promedio_anual  (solo con 2025 completo)
  - Para cats sin 2025 completo: índice desde los meses disponibles, marcado como parcial
  - Proyección 2026: aplica índice al promedio mensual 2026 observado (ene–may)
  - Hoja TOTAL muestra el agregado de las 20 categorías
  - Una hoja por categoría con gráfico de barras de estacionalidad
"""
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference, LineChart
from openpyxl.chart.series import SeriesLabel
import warnings
warnings.filterwarnings('ignore')

# ── Constantes ────────────────────────────────────────────────────────
CATS_VALIDAS = ['ACCESORIOS','ANDINA','CERDO','COLGATE','DERMODIS','DULFINA',
                'HIGIENE Y CUIDADO','HOMEPRO PERU','HUEVO','KIMBERLY','LA CORONA',
                'LA PATRONA','MEDIFARMA','PAVO','POLLO','PROCESADOS','RINTI',
                'TAMBOS PERU','VERDUM','YICHANG']
SSFF_CATS = ['CERDO','HUEVO','PAVO','POLLO','PROCESADOS']

MESES_LABEL = {1:'Ene',2:'Feb',3:'Mar',4:'Abr',5:'May',6:'Jun',
               7:'Jul',8:'Ago',9:'Sep',10:'Oct',11:'Nov',12:'Dic'}
MESES_LARGO = {1:'Enero',2:'Febrero',3:'Marzo',4:'Abril',5:'Mayo',6:'Junio',
               7:'Julio',8:'Agosto',9:'Septiembre',10:'Octubre',11:'Noviembre',12:'Diciembre'}

OUTPUT = 'C:/proyectos/SSFF/estacionalidad_ssff.xlsx'

# ── Colores ───────────────────────────────────────────────────────────
C_HDR   = '1F3864'   # azul oscuro
C_SSFF  = '1F5C2E'   # verde oscuro (cats SSFF)
C_TOTAL = '2E4057'   # azul gris (total)
C_2025  = '2E75B6'   # azul (2025)
C_2026  = 'ED7D31'   # naranja (2026)
C_PROY  = 'A9D18E'   # verde claro (proyección)
C_IDX   = 'FFF2CC'   # amarillo suave (índice)
C_IDX_H = 'F4B942'   # naranja (mes alto)
C_IDX_L = 'C9E5F5'   # azul claro (mes bajo)
C_GRIS  = 'F2F2F2'
C_BLC   = 'FFFFFF'
C_PARC  = 'FFE0CC'   # fondo naranja suave = datos parciales

def fill(hex_): return PatternFill('solid', fgColor=hex_)
def fnt(bold=False, sz=9, color='FF000000', italic=False):
    return Font(bold=bold, size=sz, color=color, italic=italic)
def aln(h='center', v='center', wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
def brd():
    s = Side(style='thin', color='BFBFBF')
    return Border(left=s, right=s, top=s, bottom=s)

# ── Carga de datos ────────────────────────────────────────────────────
print("Cargando datos...")
hist = pd.read_csv('C:/proyectos/SSFF/export_data_ssff.csv', header=None,
                   names=['PERIODO','CATEGORIA','LINEA','PEDIDOS','VOLUMEN','MONTO'])
hist['PERIODO'] = hist['PERIODO'].astype(int)
hist['anio']    = 2000 + hist['PERIODO'] // 100
hist['mes']     = hist['PERIODO'] % 100
hist = hist[hist['CATEGORIA'].isin(CATS_VALIDAS)]

# Ventas mensuales por categoría
ventas = (hist.groupby(['anio','mes','CATEGORIA'])['MONTO']
          .sum().reset_index()
          .rename(columns={'MONTO':'soles'}))

# ── Calcular índices estacionales ────────────────────────────────────
def calcular_estacional(cat):
    """
    Retorna dict con:
      - idx[mes] = índice estacional (1.0 = promedio)
      - base = 'completo' (2025 entero) o 'parcial' (meses disponibles)
      - data_2025[mes], data_2026[mes]
      - prom_2025, prom_2026_obs
    """
    df = ventas[ventas['CATEGORIA'] == cat].copy()
    data_2025 = df[df['anio']==2025].set_index('mes')['soles'].to_dict()
    data_2026 = df[df['anio']==2026].set_index('mes')['soles'].to_dict()

    # Completar con 0 si falta algún mes
    for m in range(1, 13):
        data_2025.setdefault(m, 0)
    for m in range(1, 6):
        data_2026.setdefault(m, 0)

    meses_2025_con_datos = [m for m in range(1,13) if data_2025.get(m,0) > 0]
    base = 'completo' if len(meses_2025_con_datos) >= 10 else 'parcial'

    if base == 'completo':
        # Índice sobre 2025 completo
        prom_2025 = np.mean([data_2025[m] for m in range(1,13)])
        idx = {m: (data_2025[m] / prom_2025) if prom_2025 > 0 else 1.0 for m in range(1,13)}
    else:
        # Solo meses disponibles (2025 + 2026 combinados para más muestra)
        todos = {}
        for m in meses_2025_con_datos:
            todos[m] = data_2025[m]
        for m in range(1,6):
            if data_2026.get(m,0) > 0 and m not in todos:
                todos[m] = data_2026[m]
            elif data_2026.get(m,0) > 0:
                todos[m] = (todos[m] + data_2026[m]) / 2
        prom_parcial = np.mean(list(todos.values())) if todos else 1.0
        idx = {}
        for m in range(1,13):
            if m in todos:
                idx[m] = todos[m] / prom_parcial if prom_parcial > 0 else 1.0
            else:
                idx[m] = None  # sin datos

    prom_2025 = np.mean([data_2025[m] for m in range(1,13)])
    obs_2026  = [data_2026[m] for m in range(1,6) if data_2026.get(m,0) > 0]
    prom_2026_obs = np.mean(obs_2026) if obs_2026 else prom_2025

    # Proyección 2026: promedio_2026_obs × índice
    proy_2026 = {}
    for m in range(1,13):
        if idx.get(m) is not None:
            proy_2026[m] = prom_2026_obs * idx[m]
        else:
            proy_2026[m] = None

    return {
        'idx':         idx,
        'base':        base,
        'data_2025':   data_2025,
        'data_2026':   data_2026,
        'proy_2026':   proy_2026,
        'prom_2025':   prom_2025,
        'prom_2026':   prom_2026_obs,
    }

resultados = {cat: calcular_estacional(cat) for cat in CATS_VALIDAS}

# También calcular TOTAL (suma de las 20 cats)
total_2025 = {}
total_2026 = {}
for m in range(1,13):
    total_2025[m] = sum(resultados[c]['data_2025'].get(m,0) for c in CATS_VALIDAS)
for m in range(1,6):
    total_2026[m] = sum(resultados[c]['data_2026'].get(m,0) for c in CATS_VALIDAS)
for m in range(1,6):
    total_2026.setdefault(m,0)

prom_tot_2025 = np.mean(list(total_2025.values()))
obs_tot_2026  = np.mean([total_2026[m] for m in range(1,6)])
idx_total     = {m: total_2025[m] / prom_tot_2025 if prom_tot_2025 > 0 else 1.0 for m in range(1,13)}
proy_tot_2026 = {m: obs_tot_2026 * idx_total[m] for m in range(1,13)}

print(f"Categorías procesadas: {len(resultados)}")
for cat in CATS_VALIDAS:
    r = resultados[cat]
    print(f"  {cat:22s} base={r['base']:8s}  prom_2025=S/{r['prom_2025']:>10,.0f}  prom_2026(obs)=S/{r['prom_2026']:>10,.0f}")

# ── Construir Excel ───────────────────────────────────────────────────
print("\nGenerando Excel...")
wb = Workbook()

# ════════════════════════════════════════════════════════════════════════
# HOJA RESUMEN: índices de todas las categorías en una sola vista
# ════════════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = 'RESUMEN'
ws.freeze_panes = 'C4'

# Título
ws.merge_cells('A1:Q1')
c = ws.cell(1,1, 'ESTACIONALIDAD POR CATEGORÍA — SSFF 2025–2026')
c.font = fnt(bold=True, sz=13, color='FFFFFFFF')
c.fill = fill(C_HDR); c.alignment = aln()
ws.row_dimensions[1].height = 22

ws.merge_cells('A2:Q2')
c = ws.cell(2,1, 'Índice estacional = ventas_mes / promedio_mensual  |  >1.0 mes alto  |  <1.0 mes bajo  |  Base: 2025 completo (o parcial si la categoría arrancó después)')
c.font = fnt(sz=8, italic=True, color='FF595959')
c.alignment = Alignment(horizontal='left', vertical='center')
ws.row_dimensions[2].height = 14

# Encabezados fila 3
hdrs = ['CATEGORÍA', 'TIPO', 'ENE','FEB','MAR','ABR','MAY','JUN','JUL','AGO','SEP','OCT','NOV','DIC', 'PROM 25', 'PROM 26 OBS', 'BASE']
col_bgs = [C_HDR]*2 + [C_HDR]*12 + [C_2025, C_2026, C_HDR]
for j, (h, bg) in enumerate(zip(hdrs, col_bgs), 1):
    c = ws.cell(3, j, h)
    c.font = fnt(bold=True, sz=9, color='FFFFFFFF')
    c.fill = fill(bg); c.alignment = aln(); c.border = brd()

ws.column_dimensions['A'].width = 20
ws.column_dimensions['B'].width = 8
for col in range(3, 15):
    ws.column_dimensions[get_column_letter(col)].width = 7
ws.column_dimensions[get_column_letter(15)].width = 13
ws.column_dimensions[get_column_letter(16)].width = 13
ws.column_dimensions[get_column_letter(17)].width = 9

# Fila TOTAL primero
row = 4
tipo_color = C_TOTAL
cat_label  = 'TOTAL (20 CATS)'
for j, val in enumerate([cat_label, 'TOTAL'], 1):
    c = ws.cell(row, j, val)
    c.font = fnt(bold=True, sz=9, color='FFFFFFFF'); c.fill = fill(tipo_color); c.alignment = aln(); c.border = brd()
for m, col in enumerate(range(3, 15), 1):
    idx_v = idx_total.get(m)
    c = ws.cell(row, col, round(idx_v, 3) if idx_v else '')
    _bg = C_IDX_H if idx_v and idx_v >= 1.05 else (C_IDX_L if idx_v and idx_v <= 0.95 else C_IDX)
    c.font = fnt(bold=True, sz=9); c.fill = fill(_bg); c.alignment = aln('right'); c.border = brd()
    c.number_format = '0.000'
for j, (val, bg) in enumerate([(prom_tot_2025, C_2025),(obs_tot_2026, C_2026),('completo', C_TOTAL)], 15):
    c = ws.cell(row, j, val if isinstance(val,str) else int(round(val)))
    c.font = fnt(bold=True, sz=9, color='FFFFFFFF'); c.fill = fill(bg); c.alignment = aln('right' if j<17 else 'center'); c.border = brd()
    if j < 17: c.number_format = '#,##0'
row += 1

# Separador
for j in range(1, 18):
    ws.cell(row, j).fill = fill('D9D9D9')
row += 1

# Filas por categoría
for i, cat in enumerate(CATS_VALIDAS):
    r     = resultados[cat]
    es_ssff = cat in SSFF_CATS
    tipo  = 'SSFF' if es_ssff else 'TERCERO'
    bg_cat = C_SSFF if es_ssff else C_HDR
    bg_row = 'EBF3E8' if es_ssff else (C_GRIS if i % 2 == 0 else C_BLC)

    # Col A y B
    for j, val in enumerate([cat, tipo], 1):
        c = ws.cell(row, j, val)
        c.font = fnt(bold=(j==1), sz=9, color='FFFFFFFF' if j==1 else 'FF000000')
        c.fill = fill(bg_cat if j==1 else (C_SSFF if es_ssff else 'E0E0E0'))
        c.alignment = aln('left' if j==1 else 'center'); c.border = brd()

    # Índices meses 1–12
    for m, col in enumerate(range(3,15), 1):
        idx_v = r['idx'].get(m)
        c = ws.cell(row, col)
        if idx_v is None:
            c.value = '—'
            c.font = fnt(sz=8, color='FFAAAAAA'); c.fill = fill('F0F0F0')
        else:
            c.value = round(idx_v, 3)
            _bg = C_IDX_H if idx_v >= 1.05 else (C_IDX_L if idx_v <= 0.95 else C_IDX)
            c.font = fnt(sz=9); c.fill = fill(_bg)
            c.number_format = '0.000'
        c.alignment = aln('right'); c.border = brd()

    # Promedio 2025, promedio 2026 obs, base
    bg_base = C_PARC if r['base'] == 'parcial' else 'D5E8D4'
    for j, (val, bg) in enumerate([(r['prom_2025'],'DDEEFF'),(r['prom_2026'],'FFE5CC'),(r['base'], bg_base)], 15):
        c = ws.cell(row, j, int(round(val)) if not isinstance(val,str) else val)
        c.font = fnt(sz=9, bold=(j==17))
        c.fill = fill(bg); c.alignment = aln('right' if j<17 else 'center'); c.border = brd()
        if j < 17: c.number_format = '#,##0'
    row += 1

# ════════════════════════════════════════════════════════════════════════
# HOJA PROYECCIÓN 2026: ventas reales ene–may + proyectado jun–dic
# ════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet('PROYECCIÓN 2026')
ws2.freeze_panes = 'B4'

ws2.merge_cells('A1:N1')
c = ws2.cell(1,1,'PROYECCIÓN ANUAL 2026 POR CATEGORÍA — Aplicando índice estacional sobre promedio ene–may 2026')
c.font = fnt(bold=True, sz=12, color='FFFFFFFF'); c.fill = fill(C_HDR); c.alignment = aln()
ws2.row_dimensions[1].height = 22

ws2.merge_cells('A2:N2')
c = ws2.cell(2,1,'Verde = real observado  |  Naranja claro = proyectado (índice × promedio ene–may 2026)  |  Fondo amarillo = mes con índice alto (>1.05)')
c.font = fnt(sz=8, italic=True, color='FF595959'); c.alignment = Alignment(horizontal='left',vertical='center')
ws2.row_dimensions[2].height = 14

hdrs2 = ['CATEGORÍA','ENE','FEB','MAR','ABR','MAY','JUN','JUL','AGO','SEP','OCT','NOV','DIC','TOTAL PROY']
col_bgs2 = [C_HDR] + [C_2026]*5 + [C_PROY.replace('A9','7B')]*7 + [C_HDR]
for j,(h,bg) in enumerate(zip(hdrs2,col_bgs2),1):
    c = ws2.cell(3,j,h); c.font=fnt(bold=True,sz=9,color='FFFFFFFF')
    c.fill=fill(bg); c.alignment=aln(); c.border=brd()

ws2.column_dimensions['A'].width = 22
for col in range(2,15):
    ws2.column_dimensions[get_column_letter(col)].width = 11

# Fila TOTAL
row2 = 4
real_tot = [total_2026.get(m,0) for m in range(1,6)]
proy_tot  = [proy_tot_2026.get(m,0) for m in range(6,13)]
total_anual_tot = sum(real_tot) + sum(proy_tot)

for j,val in enumerate(['TOTAL (20 CATS)']+real_tot+proy_tot+[total_anual_tot],1):
    c = ws2.cell(row2,j,val if j==1 else int(round(val)))
    c.font = fnt(bold=True,sz=9,color='FFFFFFFF'); c.fill=fill(C_TOTAL)
    c.alignment=aln('left' if j==1 else 'right'); c.border=brd()
    if j>1: c.number_format='#,##0'
row2 += 1

for j in range(1,15):
    ws2.cell(row2,j).fill=fill('D9D9D9')
row2 += 1

for i,cat in enumerate(CATS_VALIDAS):
    r = resultados[cat]
    es_ssff = cat in SSFF_CATS
    bg_cat = C_SSFF if es_ssff else C_HDR

    c = ws2.cell(row2,1,cat)
    c.font=fnt(bold=True,sz=9,color='FFFFFFFF'); c.fill=fill(bg_cat)
    c.alignment=aln('left'); c.border=brd()

    for m in range(1,6):
        val = r['data_2026'].get(m,0)
        c = ws2.cell(row2, m+1, int(round(val)))
        c.font=fnt(sz=9); c.fill=fill('D5E8D4'); c.alignment=aln('right'); c.border=brd()
        c.number_format='#,##0'

    total_proy_cat = sum(r['data_2026'].get(m,0) for m in range(1,6))
    for m in range(6,13):
        pv = r['proy_2026'].get(m)
        val = int(round(pv)) if pv is not None else 0
        total_proy_cat += val
        idx_v = r['idx'].get(m)
        _bg = 'FFE0AA' if (idx_v and idx_v >= 1.05) else 'FFF3E0'
        c = ws2.cell(row2, m+1, val)
        c.font=fnt(sz=9,italic=True); c.fill=fill(_bg); c.alignment=aln('right'); c.border=brd()
        c.number_format='#,##0'

    c = ws2.cell(row2,14, int(round(total_proy_cat)))
    c.font=fnt(bold=True,sz=9); c.fill=fill('C6EFCE'); c.alignment=aln('right'); c.border=brd()
    c.number_format='#,##0'
    row2 += 1

# ════════════════════════════════════════════════════════════════════════
# HOJA COMPARATIVO 2025 vs 2026 PROYECTADO: variación mes a mes
# ════════════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet('COMPARATIVO')
ws3.freeze_panes = 'B4'

ws3.merge_cells('A1:N1')
c = ws3.cell(1,1,'COMPARATIVO TOTAL — Real 2025 vs Proyectado 2026 (Tipo 1 aplicado)')
c.font=fnt(bold=True,sz=12,color='FFFFFFFF'); c.fill=fill(C_HDR); c.alignment=aln()
ws3.row_dimensions[1].height=22

ws3.merge_cells('A2:N2')
c = ws3.cell(2,1,'Proyectado 2026 = promedio ene–may 2026 × índice estacional 2025  |  Azul = real 2025  |  Naranja = proyectado 2026')
c.font=fnt(sz=8,italic=True,color='FF595959'); c.alignment=Alignment(horizontal='left',vertical='center')
ws3.row_dimensions[2].height=14

hdrs3=['SERIE','ENE','FEB','MAR','ABR','MAY','JUN','JUL','AGO','SEP','OCT','NOV','DIC','TOTAL']
for j,(h,bg) in enumerate(zip(hdrs3,[C_HDR]+[C_2025]*12+[C_HDR]),1):
    c=ws3.cell(3,j,h); c.font=fnt(bold=True,sz=9,color='FFFFFFFF')
    c.fill=fill(bg); c.alignment=aln(); c.border=brd()

ws3.column_dimensions['A'].width=22
for col in range(2,15):
    ws3.column_dimensions[get_column_letter(col)].width=11

row3=4
series_data = [
    ('Real 2025',       [total_2025.get(m,0) for m in range(1,13)], C_2025),
    ('Proy. 2026',      [total_2026.get(m,0) if m<=5 else proy_tot_2026.get(m,0) for m in range(1,13)], C_2026),
]
for label, vals, color in series_data:
    c=ws3.cell(row3,1,label); c.font=fnt(bold=True,sz=9,color='FFFFFFFF')
    c.fill=fill(color); c.alignment=aln('left'); c.border=brd()
    for j,v in enumerate(vals,2):
        c=ws3.cell(row3,j,int(round(v))); c.font=fnt(bold=True,sz=9,color='FFFFFFFF')
        c.fill=fill(color); c.alignment=aln('right'); c.border=brd(); c.number_format='#,##0'
    c=ws3.cell(row3,14,int(round(sum(vals)))); c.font=fnt(bold=True,sz=9,color='FFFFFFFF')
    c.fill=fill(color); c.alignment=aln('right'); c.border=brd(); c.number_format='#,##0'
    row3+=1

# Fila variación %
c=ws3.cell(row3,1,'Var % 26 vs 25'); c.font=fnt(bold=True,sz=9); c.fill=fill(C_IDX); c.alignment=aln(); c.border=brd()
v25_list = [total_2025.get(m,0) for m in range(1,13)]
v26_list = [total_2026.get(m,0) if m<=5 else proy_tot_2026.get(m,0) for m in range(1,13)]
for j,(v25,v26) in enumerate(zip(v25_list,v26_list),2):
    var = (v26/v25-1) if v25>0 else 0
    c=ws3.cell(row3,j,var); c.font=fnt(sz=9)
    _bg = '00B050' if var>=0.03 else ('FF0000' if var<-0.03 else C_IDX)
    _fc = 'FFFFFFFF' if abs(var)>=0.03 else 'FF000000'
    c.fill=fill(_bg); c.font=fnt(sz=9,color=_fc)
    c.alignment=aln('right'); c.border=brd(); c.number_format='0.0%'
tot_var = (sum(v26_list)/sum(v25_list)-1) if sum(v25_list)>0 else 0
c=ws3.cell(row3,14,tot_var); c.font=fnt(bold=True,sz=9); c.fill=fill(C_IDX)
c.alignment=aln('right'); c.border=brd(); c.number_format='0.0%'
row3+=2

# Detalle por categoría en comparativo
ws3.cell(row3,1,'DETALLE POR CATEGORÍA').font=fnt(bold=True,sz=10)
row3+=1
hdrs_det=['CATEGORÍA','TIPO','TOTAL 2025','TOTAL PROY 2026','VAR %','PROM MES 25','PROM MES 26']
bgs_det=[C_HDR,C_HDR,C_2025,C_2026,'606060',C_2025,C_2026]
for j,(h,bg) in enumerate(zip(hdrs_det,bgs_det),1):
    c=ws3.cell(row3,j,h); c.font=fnt(bold=True,sz=9,color='FFFFFFFF')
    c.fill=fill(bg); c.alignment=aln(); c.border=brd()
ws3.column_dimensions['A'].width=22
ws3.column_dimensions['B'].width=9
for col_l in ['C','D','E','F','G']:
    ws3.column_dimensions[col_l].width=16
row3+=1

for i,cat in enumerate(CATS_VALIDAS):
    r=resultados[cat]
    es_ssff=cat in SSFF_CATS
    tot25=sum(r['data_2025'].values())
    tot26_obs=sum(r['data_2026'].get(m,0) for m in range(1,6))
    tot26_proy=tot26_obs+sum(r['proy_2026'].get(m,0) or 0 for m in range(6,13))
    var=(tot26_proy/tot25-1) if tot25>0 else 0
    bg_row=C_GRIS if i%2==0 else C_BLC
    bg_cat=C_SSFF if es_ssff else C_HDR
    tipo='SSFF' if es_ssff else 'TERCERO'
    vals_row=[cat,tipo,int(round(tot25)),int(round(tot26_proy)),var,
              int(round(r['prom_2025'])),int(round(r['prom_2026']))]
    for j,val in enumerate(vals_row,1):
        c=ws3.cell(row3,j,val)
        if j==1:
            c.font=fnt(bold=True,sz=9,color='FFFFFFFF'); c.fill=fill(bg_cat)
        elif j==2:
            c.font=fnt(sz=9); c.fill=fill('EBF3E8' if es_ssff else bg_row)
        elif j==5:
            _bg='C6EFCE' if var>=0.03 else ('FFC7CE' if var<-0.03 else bg_row)
            c.font=fnt(sz=9); c.fill=fill(_bg)
            c.number_format='0.0%'
        else:
            c.font=fnt(sz=9); c.fill=fill(bg_row); c.number_format='#,##0'
        c.alignment=aln('right' if j>2 else ('left' if j==1 else 'center')); c.border=brd()
    row3+=1

# ════════════════════════════════════════════════════════════════════════
# HOJAS INDIVIDUALES POR CATEGORÍA (con gráfico)
# ════════════════════════════════════════════════════════════════════════
for cat in CATS_VALIDAS:
    r = resultados[cat]
    es_ssff = cat in SSFF_CATS
    bg_cat  = C_SSFF if es_ssff else C_HDR
    sheet_name = cat[:31]
    wsc = wb.create_sheet(sheet_name)
    wsc.freeze_panes = 'B5'

    # Título
    wsc.merge_cells('A1:N1')
    c=wsc.cell(1,1,f'ESTACIONALIDAD — {cat}  |  Base: {r["base"].upper()}  |  Tipo: {"SSFF" if es_ssff else "TERCERO"}')
    c.font=fnt(bold=True,sz=12,color='FFFFFFFF'); c.fill=fill(bg_cat); c.alignment=aln()
    wsc.row_dimensions[1].height=22

    # Encabezados col
    wsc.cell(2,1,'MES').font=fnt(bold=True,sz=9)
    hdrs_c=['MES']+[MESES_LABEL[m] for m in range(1,13)]+['TOTAL','PROMEDIO']
    for j,h in enumerate(hdrs_c,1):
        c=wsc.cell(3,j,h); c.font=fnt(bold=True,sz=9,color='FFFFFFFF')
        c.fill=fill(bg_cat); c.alignment=aln(); c.border=brd()
    wsc.column_dimensions['A'].width=18
    for col in range(2,16):
        wsc.column_dimensions[get_column_letter(col)].width=10

    # Filas de datos
    filas = [
        ('Real 2025',    [r['data_2025'].get(m,0) for m in range(1,13)], C_2025, False),
        ('Real 2026',    [r['data_2026'].get(m,0) if m<=5 else None for m in range(1,13)], C_2026, False),
        ('Proy. 2026',   [r['data_2026'].get(m,0) if m<=5 else r['proy_2026'].get(m) for m in range(1,13)], C_PROY.replace('A9','7B'), True),
        ('Índice estac.',[r['idx'].get(m) for m in range(1,13)], C_IDX, False),
    ]
    for row_i, (label, vals, color, italic) in enumerate(filas, 4):
        c=wsc.cell(row_i,1,label); c.font=fnt(bold=True,sz=9,color='FFFFFFFF' if label!='Índice estac.' else 'FF000000')
        c.fill=fill(color); c.alignment=aln('left'); c.border=brd()
        vals_num = [v for v in vals if v is not None]
        total_val = sum(vals_num) if label!='Índice estac.' else None
        prom_val  = np.mean(vals_num) if vals_num else None
        for j,val in enumerate(vals,2):
            c=wsc.cell(row_i,j)
            if val is None:
                c.value=''; c.fill=fill('F0F0F0')
            else:
                c.value=round(val,3) if label=='Índice estac.' else int(round(val))
                idx_v = r['idx'].get(j-1)
                if label=='Índice estac.':
                    _bg = C_IDX_H if val>=1.05 else (C_IDX_L if val<=0.95 else C_IDX)
                    c.fill=fill(_bg); c.number_format='0.000'
                elif label=='Real 2026' and j-1>5:
                    c.fill=fill('F0F0F0')
                else:
                    _bg=('FFE5CC' if (label=='Proy. 2026' and j-1>5 and idx_v and idx_v>=1.05)
                         else ('FFF3E0' if label=='Proy. 2026' and j-1>5 else color))
                    c.fill=fill(_bg); c.number_format='#,##0'
            c.font=fnt(sz=9,italic=italic); c.alignment=aln('right'); c.border=brd()
        # Total y promedio
        for j2,(lbl_val,fmt) in enumerate([(total_val,'#,##0'),(prom_val,'#,##0' if label!='Índice estac.' else '0.000')],14):
            c=wsc.cell(row_i,j2)
            if lbl_val is None:
                c.value=''
            else:
                c.value=int(round(lbl_val)) if fmt=='#,##0' else round(lbl_val,3)
                c.number_format=fmt
            c.font=fnt(bold=True,sz=9); c.fill=fill(color); c.alignment=aln('right'); c.border=brd()

    # Fila variación 2026 vs 2025
    row_var=8
    wsc.cell(row_var,1,'Var % 26/25').font=fnt(bold=True,sz=9); wsc.cell(row_var,1).fill=fill(C_IDX); wsc.cell(row_var,1).border=brd(); wsc.cell(row_var,1).alignment=aln()
    for m in range(1,13):
        v25=r['data_2025'].get(m,0)
        v26=r['data_2026'].get(m,0) if m<=5 else (r['proy_2026'].get(m) or 0)
        var=(v26/v25-1) if v25>0 else 0
        c=wsc.cell(row_var,m+1,var)
        _bg='C6EFCE' if var>=0.03 else ('FFC7CE' if var<-0.03 else C_IDX)
        c.font=fnt(sz=9); c.fill=fill(_bg); c.alignment=aln('right'); c.border=brd(); c.number_format='0.0%'

    # Gráfico de barras — índice estacional
    chart = BarChart()
    chart.type='col'; chart.grouping='clustered'
    chart.title=f'Índice estacional — {cat}'
    chart.y_axis.title='Índice (1.0 = promedio)'; chart.x_axis.title='Mes'
    chart.height=10; chart.width=22
    chart.y_axis.numFmt='0.00'
    chart.y_axis.scaling.min=0

    data_ref = Reference(wsc, min_col=2, max_col=13, min_row=7, max_row=7)
    cats_ref = Reference(wsc, min_col=2, max_col=13, min_row=3, max_row=3)
    chart.add_data(data_ref); chart.set_categories(cats_ref)
    chart.series[0].title = SeriesLabel(v='Índice estacional')
    chart.series[0].graphicalProperties.solidFill = C_2025
    wsc.add_chart(chart, 'A10')

# ── Guardar ───────────────────────────────────────────────────────────
wb.save(OUTPUT)
print(f"\nGuardado: {OUTPUT}")
print(f"Hojas: {[ws.title for ws in wb.worksheets]}")
