# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook('C:/proyectos/SSFF/funnel_censo_SSFF_REVISION.xlsx', data_only=True)

print('='*80)
print('ANALISIS COMPLETO DEL ARCHIVO EXCEL')
print('='*80)
print('\nHOJAS ENCONTRADAS:')
for i, sheet in enumerate(wb.sheetnames, 1):
    print(f'  {i}. {sheet}')

# ========== HOJA 1: FUNNEL_CENSO ==========
ws = wb['FUNNEL_CENSO']
print('\n\n' + '='*80)
print('HOJA 1: FUNNEL_CENSO')
print('='*80)
print(f'Dimensiones: {ws.dimensions}')
print(f'Max filas: {ws.max_row}, Max columnas: {ws.max_column}')
print(f'Celdas fusionadas: {list(ws.merged_cells)}')

print('\nCONTENIDO (TODAS LAS FILAS):')
for row_idx in range(1, ws.max_row + 1):
    row_data = []
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row_idx, col_idx)
        if cell.value is not None:
            val_repr = repr(cell.value)
            if len(val_repr) > 60:
                val_repr = val_repr[:60] + '...'
            row_data.append(f'{cell.coordinate}={val_repr}')
    if row_data or row_idx <= 20:
        if row_data:
            print(f'Fila {row_idx:3d}: {" | ".join(row_data)}')
        else:
            print(f'Fila {row_idx:3d}: (vacia)')

# ========== HOJA 2: DETALLE_CLIENTE ==========
ws = wb['DETALLE_CLIENTE']
print('\n\n' + '='*80)
print('HOJA 2: DETALLE_CLIENTE')
print('='*80)
print(f'Dimensiones: {ws.dimensions}')
print(f'Max filas: {ws.max_row}, Max columnas: {ws.max_column}')
print(f'Celdas fusionadas: {list(ws.merged_cells)}')
print(f'Freeze panes: {ws.freeze_panes}')

print('\nENCAbEZADOS (FILA 1):')
for col_idx in range(1, ws.max_column + 1):
    cell = ws.cell(1, col_idx)
    print(f'  Columna {chr(64+col_idx)}: {cell.value}')

print(f'\nPRIMERAS 5 FILAS DE DATOS (SAMPLE):')
for row_idx in range(2, min(7, ws.max_row + 1)):
    row_data = []
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row_idx, col_idx)
        val_repr = repr(cell.value)
        if len(val_repr) > 25:
            val_repr = val_repr[:25] + '...'
        row_data.append(val_repr)
    print(f'Fila {row_idx}: {" | ".join(row_data)}')

# ========== HOJA 3: ULTIMA_COMPRA_CENSO ==========
ws = wb['ULTIMA_COMPRA_CENSO']
print('\n\n' + '='*80)
print('HOJA 3: ULTIMA_COMPRA_CENSO')
print('='*80)
print(f'Dimensiones: {ws.dimensions}')
print(f'Max filas: {ws.max_row}, Max columnas: {ws.max_column}')
print(f'Freeze panes: {ws.freeze_panes}')
print(f'Celdas fusionadas: {list(ws.merged_cells)}')

print('\nENCAbEZADOS (FILA 2):')
for col_idx in range(1, ws.max_column + 1):
    cell = ws.cell(2, col_idx)
    print(f'  Columna {chr(64+col_idx)}: {cell.value}')

# ========== HOJA 4: ULTIMA_COMPRA_NO_CENSO ==========
ws = wb['ULTIMA_COMPRA_NO_CENSO']
print('\n\n' + '='*80)
print('HOJA 4: ULTIMA_COMPRA_NO_CENSO')
print('='*80)
print(f'Dimensiones: {ws.dimensions}')
print(f'Max filas: {ws.max_row}, Max columnas: {ws.max_column}')
print(f'Freeze panes: {ws.freeze_panes}')
print(f'Auto filter: {ws.auto_filter.ref if ws.auto_filter else "None"}')
print(f'Celdas fusionadas: {list(ws.merged_cells)}')

print('\nENCAbEZADOS (FILA 2):')
for col_idx in range(1, ws.max_column + 1):
    cell = ws.cell(2, col_idx)
    print(f'  Columna {chr(64+col_idx)}: {cell.value}')

print('\n\n' + '='*80)
print('FIN DEL ANALISIS')
print('='*80)
